import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_table_borders(table, color="D1D5DB", sz="4"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:left w:val="none" w:sz="0" w:space="0" w:color="auto"/>
            <w:right w:val="none" w:sz="0" w:space="0" w:color="auto"/>
            <w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideV w:val="none" w:sz="0" w:space="0" w:color="auto"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def add_header_footer(doc):
    for s in doc.sections[1:]:
        header = s.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("Secure Software Engineering – CTI Sharing Platform | Topic 29")
        hrun.font.name = "Calibri"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(100, 116, 139)

        footer = s.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        frun_left = fp.add_run("Abishek Singh P | 24CYS401\t\tInterim Report (Phases 1–10)")
        frun_left.font.name = "Calibri"
        frun_left.font.size = Pt(8.5)
        frun_left.font.color.rgb = RGBColor(100, 116, 139)

def format_table(table, col_widths, headers, rows_data):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    
    # Header row
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], "0F172A")
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=150, right=150)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs:
            run.font.name = "Calibri"
            run.font.size = Pt(9.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            
    # Data rows
    for r_idx, row_values in enumerate(rows_data):
        row_cells = table.add_row().cells
        bg_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_values):
            row_cells[c_idx].text = str(val)
            set_cell_background(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], top=100, bottom=100, left=150, right=150)
            p = row_cells[c_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.name = "Calibri"
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(30, 41, 59)
                
    for row in table.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = width

def add_placeholder_box(doc, title, subtitle, prompt_ref=None):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F1F5F9")
    set_cell_margins(cell, top=180, bottom=180, left=250, right=250)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="6" w:space="0" w:color="0284C7"/>
            <w:left w:val="single" w:sz="18" w:space="0" w:color="0284C7"/>
            <w:bottom w:val="single" w:sz="6" w:space="0" w:color="0284C7"/>
            <w:right w:val="single" w:sz="6" w:space="0" w:color="0284C7"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p.add_run(f"[{title}]\n")
    r1.font.name = "Calibri"
    r1.font.size = Pt(10)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(2, 132, 199)
    
    r2 = p.add_run(f"{subtitle}\n")
    r2.font.name = "Calibri"
    r2.font.size = Pt(9)
    r2.font.italic = True
    r2.font.color.rgb = RGBColor(71, 85, 105)
    
    if prompt_ref:
        r3 = p.add_run(f"Eraser AI Generation Prompt: See evidence/DIAGRAM_GENERATION_PROMPTS.md ({prompt_ref})")
        r3.font.name = "Calibri"
        r3.font.size = Pt(8.5)
        r3.font.color.rgb = RGBColor(100, 116, 139)
        
    doc.add_paragraph() # spacing

print("Script template ready.")
