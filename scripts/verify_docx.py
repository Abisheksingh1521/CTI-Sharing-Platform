import os
import docx

def verify_docx(filename):
    print(f"Verifying {filename}...")
    assert os.path.exists(filename), f"File {filename} does not exist!"
    doc = docx.Document(filename)
    print("1. Document opened successfully.")

    # Check sections / headings
    headings = []
    text_corpus = []
    for p in doc.paragraphs:
        text = p.text.strip()
        if text:
            text_corpus.append(text)
            if p.style.name.startswith('Heading') or text.startswith(('1.', '2.', '3.', '4.', '5.', '6.', '7.', '8.', '9.', '10.', '11.', '12.', 'TABLE OF CONTENTS')):
                headings.append(text)

    # Check table text
    table_text = []
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                table_text.append(cell.text)

    full_text = " ".join(text_corpus + table_text)

    # Requirement 2: Check Phase 1-10 sections
    expected_phases = [
        "1. Project Overview",
        "2. Phase 1 – Agile Development",
        "3. Phase 2 – Requirements Engineering",
        "4. Phase 3 – UML and Use Case Analysis",
        "5. Phase 4 – Relational Data Modeling and Data Flow",
        "6. Phase 5 – Architecture, Components, and Design Patterns",
        "7. Phase 6 – User Interface Design",
        "8. Phase 7 – Threat Modeling, STRIDE & Vulnerability Analysis",
        "9. Phase 8 – Attack Tree Decomposition",
        "10. Phase 9 – Product Backlog and Jira",
        "11. Phase 10 – Scrum Execution, Burndown & Velocity Metrics",
        "12. Unified End-to-End Traceability Matrix"
    ]
    for ep in expected_phases:
        found = any(ep in h for h in headings) or ep in full_text
        print(f"  - Section '{ep}': {'FOUND' if found else 'MISSING'}")
        assert found, f"Missing section: {ep}"

    # Requirement 3: Table of contents exists
    assert "TABLE OF CONTENTS" in full_text, "Missing Table of Contents!"
    print("3. Table of Contents exists.")

    # Requirement 4: No raw markdown syntax remains
    markdown_artifacts = ["```", "## ", "### ", "**A01", "| --- |", "|:---:|"]
    for md in markdown_artifacts:
        count = full_text.count(md)
        print(f"  - Markdown artifact '{md}': count = {count}")
        assert count == 0, f"Found raw markdown syntax '{md}' in document!"
    print("4. Zero raw Markdown syntax detected.")

    # Requirement 5: Diagram placeholders clearly identified
    placeholders = [
        "DIAGRAM: UML USE CASE DIAGRAM",
        "DIAGRAM: RELATIONAL ENTITY-RELATIONSHIP (ER) DIAGRAM",
        "DIAGRAM: LEVEL 0 CONTEXT DATA FLOW DIAGRAM",
        "DIAGRAM: LEVEL 1 DATA FLOW DIAGRAM WITH TRUST BOUNDARIES",
        "DIAGRAM: LAYERED HEXAGONAL SOFTWARE & SECURITY ARCHITECTURE",
        "INSERT ACTUAL APPLICATION SCREENSHOT: SCREEN 1",
        "INSERT ACTUAL APPLICATION SCREENSHOT: SCREEN 2",
        "INSERT ACTUAL APPLICATION SCREENSHOT: SCREEN 3",
        "INSERT ACTUAL APPLICATION SCREENSHOT: SCREEN 4",
        "DIAGRAM: PHASE 7 STRIDE THREAT MODEL",
        "DIAGRAM: PHASE 8 ATTACK TREE",
        "INSERT ACTUAL JIRA SCREENSHOT",
        "INSERT ACTUAL JIRA METRIC SCREENSHOT"
    ]
    for ph in placeholders:
        found = ph in full_text
        print(f"  - Placeholder '{ph}': {'FOUND' if found else 'MISSING'}")
        assert found, f"Missing placeholder: {ph}"
    print("5. All diagram and screenshot placeholders clearly identified.")

    # Check forbidden terminology
    forbidden = ["blockchain", "non-repudiation", "non-forgeable"]
    for fb in forbidden:
        assert fb not in full_text.lower(), f"Forbidden terminology '{fb}' found in document!"
    print("6. Zero forbidden terminology detected ('blockchain', 'non-repudiation', 'non-forgeable').")

    # Requirement 7: Check text colors across all paragraphs and tables
    non_black_runs = []
    for p in doc.paragraphs:
        for r in p.runs:
            if r.font.color and r.font.color.rgb:
                rgb_hex = str(r.font.color.rgb)
                if rgb_hex != "000000":
                    non_black_runs.append((r.text[:30], rgb_hex))

    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        if r.font.color and r.font.color.rgb:
                            rgb_hex = str(r.font.color.rgb)
                            if rgb_hex != "000000":
                                non_black_runs.append((r.text[:30], rgb_hex))

    print(f"7. Non-black text runs count: {len(non_black_runs)}")
    assert len(non_black_runs) == 0, f"Found non-black text runs: {non_black_runs[:5]}"
    print("   -> All explicit font colors are 100% pure black (#000000).")

    print("\nALL VERIFICATION CHECKS PASSED PERFECTLY!")

if __name__ == '__main__':
    verify_docx("CTI_Sharing_Platform_Phase_01_to_10.docx")
    verify_docx("CTI_Sharing_Platform_Phase_01_to_10_Updated.docx")
